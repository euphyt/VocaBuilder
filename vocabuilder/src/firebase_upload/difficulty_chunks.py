import json
from openpyxl import load_workbook
import os

# =========================================
# CONFIG - Manually set ranges for each chunk
CHUNK_RANGES = [
    (0, 17299),       # chunk 1: indices inclusive
    (17300, 34894),   # chunk 2
    (34895, 52490)    # chunk 3
]

# Use script folder for file paths
script_dir = os.path.dirname(os.path.abspath(__file__))
INPUT_FILE = os.path.join(script_dir, "SUBTLEX_wordnet_checked.xlsx")
# =========================================

def create_chunks_from_xlsx(filepath, chunk_ranges):
    workbook = load_workbook(filename=filepath)
    sheet = workbook.active

    # Read first column into list
    words = []
    for row in sheet.iter_rows(values_only=True):
        if row[0] is not None:
            words.append(str(row[0]).strip())

    total_words = len(words)
    print(f"Total words available: {total_words}\n")

    # Process each chunk range
    for index, (start, end) in enumerate(chunk_ranges):
        print(f"Creating chunk {index + 1}: range {start}–{end}")

        json_object = {}

        # Fill JSON only for indices that exist
        for i in range(start, end + 1):
            if i < total_words:
                json_object[i] = words[i]

        # Save JSON files in the script folder
        if index == 0:
            filename = os.path.join(script_dir, "advanced.json")
        elif index == 1:
            filename = os.path.join(script_dir, "intermediate.json")
        elif index == 2:
            filename = os.path.join(script_dir, "beginner.json")

        with open(filename, "w", encoding="utf-8") as f:
            json.dump(json_object, f, indent=2, ensure_ascii=False)

        print(f"Saved {filename} with {len(json_object)} words\n")


if __name__ == "__main__":
    create_chunks_from_xlsx(INPUT_FILE, CHUNK_RANGES)

